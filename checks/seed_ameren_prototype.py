"""Seed the Ameren prototype from the workbook's already-issued invoices.

READS A WORKBOOK PATH, COMMITS NOTHING. The source file holds 555 real invoice lines
naming 161 real workers for a named client. It is passed in with --workbook and is
never written into this repository, which 30-odd people can read. The code is the
deliverable here; the data stays where it is.

WHAT IT PRODUCES: the gold export rows, and the issuance-ledger rows for invoices that
were issued on 2026-07-29 before any of this existed. Nothing is written to a
workspace -- deploying the rows is a separate step with its own target and its own
profile, and this file deliberately cannot do it.

WHY IT GATES BEFORE IT EMITS. The sheet is a customer-facing document that has already
been sent. If AME003, AME006 or AME008 do not hold over it, the right conclusion is
that this file is not what we think it is -- not that we should seed from it anyway.
The gates run first and refuse the whole file, because a partially-seeded gold table is
the failure this design exists to avoid.
"""
from __future__ import annotations

# DEF-12: serverless `spark_python_task` exec()s this file and does NOT define
# __file__, so every Path(__file__) below raised NameError and the file died before
# its own logic ran. compile() still records the real path in the code object.
if "__file__" not in globals():  # noqa: F821
    import inspect as _inspect

    __file__ = _inspect.currentframe().f_code.co_filename

import argparse
import collections
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "checks"))

from accelerator.ameren_prototype import (  # noqa: E402
    GOLD_FIELDS, SHEET_HEADERS, gold_row, issuance_rows,
)
from accelerator.invoice_rules import reconciles  # noqa: E402

DEFAULT_SHEET = "3 Ameren Custom File"
#: The sheet's headers sit on row 3 -- row 1 is a title and row 2 is blank.
HEADER_ROW = 3


def read_sheet(workbook: str, sheet: str) -> list[dict]:
    """Raw rows as header->value dicts.

    openpyxl is imported HERE, not at module scope, and is an optional dependency --
    the same pattern conformance_check.py uses for databricks-sdk. The offline suites
    import this module to test the pure functions around it and must not need a
    spreadsheet library to do so.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised by the message, not the path
        raise SystemExit(
            "openpyxl is required to read the workbook: "
            "uv run --frozen --with openpyxl python checks/seed_ameren_prototype.py ..."
        ) from exc

    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not in {wb.sheetnames}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = [("" if c is None else str(c).strip()) for c in rows[HEADER_ROW - 1]]
    missing = [h for h in SHEET_HEADERS if h not in header]
    if missing:
        raise ValueError(f"sheet {sheet!r} is missing column(s) {missing}; found {header}")
    idx = {h: header.index(h) for h in SHEET_HEADERS}
    out = []
    for r in rows[HEADER_ROW:]:
        if not r or not any(c is not None and str(c).strip() for c in r):
            continue
        out.append({h: r[i] for h, i in idx.items()})
    return out


def gate(rows: list[dict]) -> list[str]:
    """Problems that stop the whole file being seeded. Empty means it may be.

    AME008 per invoice, AME003 (one gross per invoice) and AME006 (a complete 1..n line
    sequence) -- the same three the design gates a release on, applied to a document
    that has already been released.
    """
    problems: list[str] = []
    by_invoice: dict[str, list[dict]] = collections.defaultdict(list)
    for r in rows:
        by_invoice[r["invoice_number"]].append(r)

    for invoice, lines in sorted(by_invoice.items()):
        grosses = {l["invoice_amount"] for l in lines}
        if len(grosses) > 1:
            problems.append(f"AME003 {invoice}: {len(grosses)} different invoice totals {sorted(grosses)}")
            continue
        gross = next(iter(grosses))
        if not reconciles([l["amount"] for l in lines], gross):
            total = sum(l["amount"] for l in lines)
            problems.append(f"AME008 {invoice}: lines sum to {total}, gross is {gross}")
        numbers = sorted(l["line_number"] for l in lines)
        if numbers != list(range(1, len(numbers) + 1)):
            problems.append(f"AME006 {invoice}: line numbers {numbers} are not a complete 1..n sequence")
    return problems


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", required=True,
                    help="path to the Ameren workbook; NEVER commit this file")
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--run-id", default="prototype-seed",
                    help="recorded as issued_by_run_id on every ledger row")
    ap.add_argument("--out-dir", help="write gold.ndjson and issuance.ndjson here")
    args = ap.parse_args()

    raw = read_sheet(args.workbook, args.sheet)
    rows = [gold_row(r) for r in raw]

    problems = gate(rows)
    if problems:
        print(f"SEED REFUSED -- {len(problems)} problem(s):")
        for p in problems[:20]:
            print(f"  {p}")
        return 1

    ledger = issuance_rows(rows, issued_by_run_id=args.run_id)
    invoices = len({r["invoice_number"] for r in rows})
    print(f"gold rows      : {len(rows)} across {invoices} invoice(s)")
    print(f"issuance rows  : {len(ledger)}")
    print(f"issued date(s) : {sorted({r['invoice_date'] for r in ledger})}")
    print("AME003/AME006/AME008 all hold over this file.")

    if args.out_dir:
        out = Path(args.out_dir)
        out.mkdir(parents=True, exist_ok=True)
        (out / "gold.ndjson").write_text(
            "\n".join(json.dumps({k: str(v) for k, v in r.items()}) for r in rows) + "\n")
        (out / "issuance.ndjson").write_text(
            "\n".join(json.dumps({k: str(v) for k, v in r.items()}) for r in ledger) + "\n")
        print(f"wrote {out}/gold.ndjson and {out}/issuance.ndjson")
    return 0


if __name__ == "__main__":
    _rc = main()
    if _rc:
        sys.exit(_rc)
